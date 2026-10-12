from pathlib import Path
import zipfile, csv, io, json, re, sys, pickle, textwrap, shutil, hashlib
from collections import defaultdict
import xml.etree.ElementTree as ET
import numpy as np
import pandas as pd
import openpyxl
from scipy.stats import hypergeom, ttest_ind
from PIL import Image
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import Circle, FancyArrowPatch
import pymupdf as fitz
sys.path.insert(0,str(Path.cwd()/'tmp/pdfs/python_libs'))
import xlrd

ROOT=Path.cwd(); SRC=ROOT/'2026 崔珣 合成致死课题原始数据'; TMP=ROOT/'tmp/pdfs'
OUT=ROOT/'output/pdf/Fig4_Fig5_S9_S12'; OUT.mkdir(parents=True,exist_ok=True)
DATA=OUT/'data';DATA.mkdir(exist_ok=True)
CELLS=['HT29','SW620']; NC='#777D83';KD='#8A3048';BLUE='#3978A4';COLORS=[NC,KD]
plt.rcParams.update({'font.family':'Arial','font.size':7,'pdf.fonttype':42,'svg.fonttype':'none','axes.linewidth':.6,'xtick.major.size':2,'ytick.major.size':2,'axes.labelpad':3,'savefig.facecolor':'white'})
PROVENANCE={}; STATS=[]
def bh(p):
    p=np.asarray(p);ix=np.argsort(p);q=np.empty(len(p));q[ix]=np.minimum.accumulate((p[ix]*len(p)/np.arange(1,len(p)+1))[::-1])[::-1];return np.minimum(q,1)
def readtsv(z,n):return pd.read_csv(io.BytesIO(z.read(n)),sep='\t')
def source(p):return str(p.relative_to(ROOT))
def extract():
    cache=TMP/'fig45_data.pkl'
    if cache.exists():return pickle.loads(cache.read_bytes())
    deg={};expr={};pca={};corr={};go_original={};gsea={};maps={};terms={};provenance={}
    for c in CELLS:
        zp=next((ROOT/'测序数据'/(c+' RNA测序数据')).glob('03.Result*.zip'));provenance[c]=source(zp)
        with zipfile.ZipFile(zp) as z:
            names=z.namelist();prefix=next(n.split('/4.Differential/')[0] for n in names if '/Result_' in n and '/4.Differential/' in n)
            deg[c]=readtsv(z,next(n for n in names if n.startswith(prefix) and n.endswith('/shTCF7L2vsshNC_deg.xls'))).set_index('gene_id')
            expr[c]=readtsv(z,prefix+'/3.Quant/1.Count/gene_fpkm.xls').set_index('gene_id')
            pca[c]=readtsv(z,prefix+'/3.Quant/2.Correlation/PCA2D.xls')
            corr[c]=readtsv(z,prefix+'/3.Quant/2.Correlation/correlation.xls')
            go_original[c]=readtsv(z,next(n for n in names if n.startswith(prefix) and n.endswith('/ALL/shTCF7L2vsshNC.all_GOenrich.xls')))
            gp=prefix+'/5.Enrichment/1.GO/shTCF7L2vsshNC/GSEA/edb/'
            edb=ET.fromstring(z.read(gp+'results.edb'))
            gsea[c]=pd.DataFrame([{k:x.get(k) for k in ['GENESET','ES','NES','NP','FDR','FWER','HIT_INDICES','ES_PROFILE']} for x in edb.iter('DTG')])
            for col in ['ES','NES','NP','FDR','FWER']:gsea[c][col]=pd.to_numeric(gsea[c][col])
            if c=='HT29':
                # Use the delivered GO annotation, including ancestors; do not download or change versions.
                with z.open(prefix+'/0.SupFile/go.xls') as raw:
                    for row in csv.DictReader(io.TextIOWrapper(raw,encoding='utf-8-sig'),delimiter='\t'):
                        gid=row['go_id'];maps.setdefault(gid,set()).add(row['gene_id']);terms[gid]=(row['go_term'],row['go_ontology'])
                provenance['GO_annotation']=prefix+'/0.SupFile/go.xls (HT29 delivered archive)'
    up={c:set(d.index[(d.padj<=.05)&(d.log2FoldChange>=1)]) for c,d in deg.items()}
    down={c:set(d.index[(d.padj<=.05)&(d.log2FoldChange<=-1)]) for c,d in deg.items()}
    common_up=up['HT29']&up['SW620'];common_down=down['HT29']&down['SW620']
    # Universe: genes with an adjusted P value in both analyses AND a BP annotation.
    tested=set(deg['HT29'].index[deg['HT29'].padj.notna()])&set(deg['SW620'].index[deg['SW620'].padj.notna()])
    bpgenes=set().union(*(s for g,s in maps.items() if terms[g][1]=='BP'))
    universe=tested&bpgenes
    enrich={}
    for direction,query in [('up',common_up),('down',common_down)]:
        qset=query&universe;rows=[]
        for gid,members in maps.items():
            if terms[gid][1]!='BP':continue
            bg=members&universe
            if not 10<=len(bg)<=500:continue
            overlap=qset&bg;k=len(overlap);N=len(universe);n=len(qset);M=len(bg)
            rows.append([gid,terms[gid][0],k,n,M,N,float(hypergeom.sf(k-1,N,M,n)), '/'.join(sorted(overlap))])
        df=pd.DataFrame(rows,columns=['GOID','Description','Count','QueryAnnotated','TermBackground','Universe','pvalue','geneIDs']);df['padj']=bh(df.pvalue);df['GeneRatio']=df.Count/df.QueryAnnotated;enrich[direction]=df.sort_values(['padj','pvalue','GOID'])
    obj=dict(deg=deg,expr=expr,pca=pca,corr=corr,go_original=go_original,gsea=gsea,maps=maps,terms=terms,up=up,down=down,common_up=common_up,common_down=common_down,universe=universe,enrich=enrich,provenance=provenance)
    cache.write_bytes(pickle.dumps(obj));return obj

D=extract();PROVENANCE.update(D['provenance'])
for c in CELLS:
    d=D['deg'][c];d.to_csv(DATA/f'{c}_DEG_full.csv',encoding='utf-8-sig');D['go_original'][c].to_csv(DATA/f'{c}_GO_original.csv',index=False)
    D['pca'][c].to_csv(DATA/f'{c}_PCA_source_coordinates.csv',index=False);D['corr'][c].to_csv(DATA/f'{c}_correlation_source.csv',index=False)
    D['gsea'][c].drop(columns=['HIT_INDICES','ES_PROFILE']).to_csv(DATA/f'{c}_GSEA_original.csv',index=False)
for direction in ['up','down']:
    D['enrich'][direction].to_csv(DATA/f'shared_{direction}_GO_recomputed.csv',index=False)
    ids=sorted(D['common_'+direction]);df=pd.DataFrame({'gene_id':ids,'gene_name':[D['deg']['HT29'].loc[g,'gene_name'] for g in ids]});df.to_csv(DATA/f'shared_{direction}_genes.csv',index=False)
print('COUNTS', {c:(len(D['up'][c]),len(D['down'][c])) for c in CELLS},'shared',len(D['common_up']),len(D['common_down']),'universe',len(D['universe']),flush=True)
for k,df in D['enrich'].items():
    print(k,'sig',sum(df.padj<.05),df[['Description','Count','QueryAnnotated','padj']].head(8).to_string(index=False),flush=True)

# Animal records: preserve IDs in each assay separately.
animalpath=next(SRC.rglob('HT29接瘤数据第二批原始数据.xlsx'));s=openpyxl.load_workbook(animalpath,data_only=True).active
animal={};days=np.array([4,7,9,11,14,16,18,20,22]);bodydays=np.r_[0,days]
for gi,g in enumerate(['shNC','shTCF7L2']):
    rr=list(range(3+gi*8,11+gi*8));animal[g]={'id':[s.cell(r,2).value for r in rr],'volume':np.array([[s.cell(r,c).value for c in range(7,40,4)] for r in rr],float),'weight':np.array([s.cell(r,40).value for r in rr],float),'body':np.array([[s.cell(r,c).value for c in [3]+list(range(4,37,4))] for r in rr],float)}
    for r in rr:
        for c in range(7,40,4):
            a,b,v=s.cell(r,c-2).value,s.cell(r,c-1).value,s.cell(r,c).value
            if a is not None and b is not None:assert np.isclose(a*b*b/2,v)
rows=[]
for g,a in animal.items():
    for i,id in enumerate(a['id']):
        rows.append([g,id,0,None,a['body'][i,0],a['weight'][i]])
        for j,t in enumerate(days):rows.append([g,id,t,a['volume'][i,j],a['body'][i,j+1],a['weight'][i]])
pd.DataFrame(rows,columns=['group','animal_id_as_recorded','day','tumor_volume_mm3','body_weight_g','day23_tumor_weight_g']).to_csv(DATA/'animal_measurements.csv',index=False)
PROVENANCE['animal']=source(animalpath)
tpath=next(SRC.rglob('TUNEL分析.xlsx'));s=openpyxl.load_workbook(tpath,data_only=True).active;tunel={};rows=[]
for g,starts in [('shNC',range(2,26,3)),('shTCF7L2',range(28,52,3))]:
    ids=[];areas=[];ints=[]
    for r in starts:
        id=s.cell(r,1).value;vals=[s.cell(r+j,7).value for j in range(3)];iv=[s.cell(r+j,6).value for j in range(3)];ids.append(id);areas.append(vals);ints.append(iv)
        for j in range(3):rows.append([g,id,j+1,vals[j],iv[j]])
    tunel[g]={'id':ids,'fields':np.array(areas,float),'area':np.mean(areas,axis=1),'IntDen':np.mean(ints,axis=1)}
pd.DataFrame(rows,columns=['group','tissue_id_as_recorded','field','positive_area_percent','integrated_density']).to_csv(DATA/'TUNEL_fields.csv',index=False)
PROVENANCE['TUNEL']=source(tpath)
# Pair the two wells of each sample; do not use six technical wells as biological n.
qpath=next(SRC.rglob('HT29 shTCF7L2 细胞 qpcr.pzfx'));ns={'p':'htt<LOCAL_PATH_REMOVED>'};qpcr={}
for t in ET.parse(qpath).getroot().findall('p:Table',ns):
    gene=t.findtext('p:Title',namespaces=ns);a=np.array([[float(v.text) for v in y.findall('.//p:d',ns)] for y in t.findall('p:YColumn',ns)])
    assert a.shape==(2,6);qpcr[gene]=a.reshape(2,3,2).mean(axis=2)
pd.DataFrame([[gene,g,i+1,v] for gene,a in qpcr.items() for g,vs in zip(['shNC','shTCF7L2'],a) for i,v in enumerate(vs)],columns=['gene','group','sample','mean_of_two_cached_normalized_wells']).to_csv(DATA/'HT29_qPCR_sample_means.csv',index=False)
PROVENANCE['qPCR']=source(qpath)

if '--data-only' in sys.argv:sys.exit(0)

FIGURES={};HEATMAP_ROWS=[]
def fig(w=210,h=230):return plt.figure(figsize=(w/25.4,h/25.4))
def text(f,x,y,s,**kw):return f.text(x,y,s,**kw)
def panel(f,x,y,letter,title):
    text(f,x,y,letter,fontsize=11,fontweight='bold');text(f,x+.033,y,title,fontsize=8,fontweight='bold')
def legend(f,x=.76,y=.978):
    text(f,x,y,'● shNC',color=NC,fontsize=7);text(f,x+.105,y,'● shTCF7L2',color=KD,fontsize=7)
def clean(ax):ax.spines[['top','right']].set_visible(False);ax.tick_params(labelsize=7)
def footer(f,s):text(f,.04,.018,s,fontsize=6,color='#454545')
def addfig(name,f):FIGURES.setdefault(name,[]).append(f)
def stars(p):return '****' if p<.0001 else '***' if p<.001 else '**' if p<.01 else '*' if p<.05 else 'ns'
def dots(ax,arrays,ylabel,labels=('shNC','shTCF7L2'),test=None):
    vmax=max(np.max(v) for v in arrays)
    for i,v in enumerate(arrays):
        v=np.asarray(v);ax.scatter(i+np.linspace(-.10,.10,len(v)),v,s=18,c=COLORS[i],edgecolor='white',lw=.35,zorder=3)
        ax.errorbar(i,np.mean(v),yerr=np.std(v,ddof=1)/np.sqrt(len(v)),fmt='_',color='black',capsize=4,ms=17,lw=.8,zorder=4)
    ax.set(xticks=[0,1],xticklabels=labels,ylabel=ylabel,xlim=(-.45,1.45),ylim=(0,vmax*1.28));clean(ax)
    if test:
        p=float(ttest_ind(*arrays,equal_var=False).pvalue);STATS.append(dict(metric=test,test='two-sided Welch t test',pvalue=p,n1=len(arrays[0]),n2=len(arrays[1])))
        y=vmax*1.10;ax.plot([0,0,1,1],[y,y+vmax*.03,y+vmax*.03,y],c='black',lw=.6);ax.text(.5,y+vmax*.05,stars(p) if p<.05 else f'P = {p:.3f}',ha='center',fontsize=8)
def imageax(f,box,p,crop=None,title=None):
    im=Image.open(p);im=im.crop(crop) if crop else im
    ax=f.add_axes(box);ax.imshow(im);ax.axis('off')
    if title:ax.set_title(title,fontsize=7,pad=4)
    return ax
def growth(ax,key='volume',individual=False):
    times=days if key=='volume' else bodydays
    for gi,(g,a) in enumerate(animal.items()):
        vs=a[key]
        if individual:
            for id,v in zip(a['id'],vs):ax.plot(times,v,lw=.7,color=COLORS[gi],alpha=.65)
        else:ax.errorbar(times,vs.mean(axis=0),yerr=vs.std(axis=0,ddof=1)/np.sqrt(8),color=COLORS[gi],fmt='o-',ms=3,capsize=2,lw=1)
    ax.set(xlabel='Days after implantation',ylabel='Tumor volume (mm³)' if key=='volume' else 'Body weight (g)',xticks=[0,7,14,22],xlim=(0,23));clean(ax)
    if key=='volume':ax.set_ylim(bottom=0)
def select_tunel(g,id,field=1):
    root=SRC/'Figure 4/Figure 4-4/2 皮下瘤组织 免疫荧光/2 IF拍摄原始照片'/g
    candidates=[p for p in root.glob('*.jpg') if re.search(rf'\s{id}肿-Merge_10\.00X-{field}\.jpg$',p.name,re.I)]
    if len(candidates)!=1:raise ValueError((g,id,field,candidates))
    return candidates[0]
photo=SRC/'Figure 3/Figure 3-3 动物实验/1 解剖结果/解剖肿瘤.jpg'

# Fig4. No Ki-67 quantification is mixed in because its sample labels are unresolved.
f=fig(180,195);legend(f,.71,.98)
panel(f,.035,.94,'A','HT29 xenograft design')
steps=[('HT29 cells','shNC / shTCF7L2'),('Subcutaneous implantation','Male BALB/c nude mice'),('Growth monitoring','8 mice / group'),('Endpoint analysis','Day 23')]
for i,(a,b) in enumerate(steps):
    y=.872-i*.046;text(f,.07,y,a,fontsize=7,fontweight='bold');text(f,.07,y-.018,b,fontsize=6.5,color='#555555')
    if i<3:text(f,.39,y-.03,'↓',fontsize=11,color=NC)
panel(f,.49,.94,'B','Excised tumors')
imageax(f,[.52,.70,.44,.22],photo,crop=(75,510,1040,930))
text(f,.52,.69,'Upper: shNC | Lower: shTCF7L2 (8 tumors each)',fontsize=6)
panel(f,.035,.64,'C','Tumor growth');panel(f,.54,.64,'D','Endpoint tumor weight')
growth(f.add_axes([.12,.40,.34,.205]))
dots(f.add_axes([.64,.40,.29,.205]),[a['weight'] for a in animal.values()],'Tumor weight (g)',test='Fig4D tumor weight, day23')
panel(f,.035,.32,'E','TUNEL / DAPI');panel(f,.69,.32,'F','TUNEL-positive area')
representatives={}
for gi,(g,a) in enumerate(tunel.items()):
    order=np.argsort(a['area']);ix=order[(len(order)-1)//2];id=a['id'][ix];field=int(np.argmin(abs(a['fields'][ix]-a['area'][ix])))+1;p=select_tunel(g,id,field);representatives[g]=(id,field,p)
    imageax(f,[.045+gi*.31,.11,.295,.17],p,title=g+f' | tissue {id}')
dots(f.add_axes([.77,.11,.18,.16]),[a['area'] for a in tunel.values()],'Positive area (%)',test='Fig4F TUNEL area, sample means')
text(f,.045,.065,'Mean ± SEM; TUNEL: 3 fields averaged per tissue. Native image scale bars retained.',fontsize=6)
text(f,.045,.044,'Endpoint comparisons: two-sided Welch tests; unadjusted P values.',fontsize=6)
footer(f,'Growth and tissue IDs are preserved separately; the shTCF7L2 9 / 36 discrepancy remains unresolved.')
addfig('Fig4',f)

# S9: individual measurements and additional independent tissue images.
f=fig(210,250);legend(f)
for gi,(g,a) in enumerate(animal.items()):
    panel(f,.04+gi*.50,.925,chr(65+gi),g+' | Individual tumor growth')
    ax=f.add_axes([.115+gi*.50,.665,.34,.225])
    for i,(id,v) in enumerate(zip(a['id'],a['volume'])):
        ax.plot(days,v,lw=.8,color=plt.get_cmap('tab10')(i),label=str(id))
    ax.set(xlabel='Days after implantation',ylabel='Tumor volume (mm³)',ylim=(0,None),xticks=[4,9,14,18,22]);clean(ax)
    ax.legend(title='Recorded ID',fontsize=5.5,title_fontsize=6,ncol=2,loc='upper left',frameon=False)
panel(f,.04,.58,'C','Endpoint tumor volume');panel(f,.54,.58,'D','Body weight')
dots(f.add_axes([.115,.335,.34,.21]),[a['volume'][:,-1] for a in animal.values()],'Tumor volume, day 22 (mm³)',test='S9C endpoint volume, day22')
growth(f.add_axes([.615,.335,.34,.21]),'body')
panel(f,.04,.255,'E','Full photograph with native ruler')
imageax(f,[.075,.035,.45,.195],photo)
text(f,.58,.18,'Upper row: shNC, 8 tumors\nLower row: shTCF7L2, 8 tumors\n\nTumor positions are not assigned animal IDs\nwithout a verified photo-to-record key.',fontsize=8,linespacing=1.6)
footer(f,'Source blank length/width entries with cached volume 0 are retained as recorded. No safety or toxicity conclusion is inferred.')
addfig('FigS9',f)
f=fig(210,205);panel(f,.04,.94,'F','Additional TUNEL / DAPI fields')
for gi,(g,a) in enumerate(tunel.items()):
    chosen=[a['id'][np.argsort(a['area'])[k]] for k in [1,3,6]]
    for j,id in enumerate(chosen):
        ix=a['id'].index(id);field=int(np.argmin(abs(a['fields'][ix]-a['area'][ix])))+1
        imageax(f,[.045+j*.32,.59-gi*.31,.29,.245],select_tunel(g,id,field),title=f'{g} | tissue {id} | field {field}')
text(f,.06,.15,'Fields are selected from three different tissues per group, spanning the measured area distribution.\nThe field closest to each tissue mean is displayed. Original scale bars and image intensities are retained.\nTissue IDs: shNC 5, 17, 24, 26, 30, 34, 38, 39; shTCF7L2 8, 12, 15, 23, 31, 33, 36, 42.',fontsize=8,linespacing=1.6)
footer(f,'These are representative fields; the statistical unit in Fig4F is the tissue-level mean, not an individual field.')
addfig('FigS9',f)

# Generic vector heatmap. Z scores are calculated separately within each cell line.
SAMPLECOLS=['shNC_1','shNC_2','shNC_3','shTCF7L2_1','shTCF7L2_2','shTCF7L2_3']
def matrix(ids,cells=CELLS):
    values=[]
    for c in cells:
        a=np.log2(D['expr'][c].reindex(ids)[SAMPLECOLS].to_numpy(float)+1)
        assert np.isfinite(a).all()
        sd=a.std(axis=1,keepdims=True);sd[sd==0]=1;values.append((a-a.mean(axis=1,keepdims=True))/sd)
    return np.hstack(values)
def heatmap(f,box,ids,cells=CELLS,labelsize=6,title=None,export=None,colorbar=True):
    ax=f.add_axes(box);a=matrix(ids,cells);im=ax.pcolormesh(a,cmap='RdBu_r',vmin=-2,vmax=2,edgecolors='none',rasterized=False)
    names=[str(D['expr'][cells[0]].loc[g,'gene_name']) for g in ids]
    ax.set(yticks=np.arange(len(ids))+.5,yticklabels=names,xticks=np.arange(6*len(cells))+.5,xticklabels=['NC1','NC2','NC3','KD1','KD2','KD3']*len(cells),xlim=(0,6*len(cells)),ylim=(len(ids),0))
    ax.tick_params(length=0,labelsize=labelsize,pad=2);ax.tick_params(axis='x',labelrotation=90)
    for j,c in enumerate(cells):
        ax.text((j+.5)/len(cells),1.035,c,ha='center',transform=ax.transAxes,fontsize=7,fontweight='bold')
        ax.plot([j*6+.08,j*6+2.92],[-.35,-.35],color=NC,lw=3,clip_on=False);ax.plot([j*6+3.08,j*6+5.92],[-.35,-.35],color=KD,lw=3,clip_on=False)
    if len(cells)==2:ax.axvline(6,color='white',lw=2)
    if colorbar:
        cb=f.colorbar(im,ax=ax,fraction=.035,pad=.03);cb.solids.set_rasterized(False);cb.ax.tick_params(labelsize=6,length=2);cb.set_label('Within-cell-line z score',fontsize=6)
    if export:
        df=pd.DataFrame(a,index=ids,columns=[c+'_'+s for c in cells for s in SAMPLECOLS]);df.insert(0,'gene_name',names);df.to_csv(DATA/(export+'.csv'))
    return ax
def dotenrich(f,box,df,fontsize=6.5):
    ax=f.add_axes(box);df=df.copy().iloc[::-1];v=-np.log10(np.maximum(df.padj.to_numpy(float),1e-300));x=df.GeneRatio.to_numpy(float)
    size=18+df.Count.to_numpy(float)*4
    sc=ax.scatter(x,np.arange(len(df)),s=size,c=v,cmap='viridis',edgecolors=['black' if q<.05 else '#B9B9B9' for q in df.padj],linewidths=.7,vmin=0,vmax=max(1.4,v.max()))
    labels=['\n'.join(textwrap.wrap(s,43)) for s in df.Description]
    ax.set(yticks=np.arange(len(df)),yticklabels=labels,xlabel='Gene ratio',ylim=(-.7,len(df)-.3));ax.tick_params(axis='y',length=0,labelsize=fontsize);ax.tick_params(axis='x',labelsize=6);ax.spines[['top','right','left']].set_visible(False)
    ax.margins(x=.15)
    cb=f.colorbar(sc,ax=ax,fraction=.045,pad=.025);cb.solids.set_rasterized(False);cb.ax.tick_params(labelsize=6);cb.set_label('-log10(FDR)',fontsize=6)
    return ax
def volcano(ax,c):
    d=D['deg'][c];x=d.log2FoldChange.to_numpy();p=d.padj.to_numpy();valid=np.isfinite(x)&np.isfinite(p);x=x[valid];p=p[valid];names=d.gene_name.to_numpy()[valid]
    # Cap for display only. Exact adjusted P values remain in the exported table.
    y=np.minimum(-np.log10(np.maximum(p,1e-300)),60);up=(p<=.05)&(x>=1);down=(p<=.05)&(x<=-1)
    for m,color in [(~(up|down),'#C9CDD0'),(down,BLUE),(up,KD)]:ax.scatter(x[m],y[m],s=3,c=color,alpha=.7,linewidths=0)
    ax.axvline(1,c='#999999',ls='--',lw=.5);ax.axvline(-1,c='#999999',ls='--',lw=.5);ax.axhline(-np.log10(.05),c='#999999',ls='--',lw=.5)
    candidates=['CXCL10','CXCL11','IFNB1'] if c=='HT29' else ['HLA-B','PSMB8','PSMB9']
    for j,g in enumerate(candidates):
        ind=np.where((names==g)&up)[0]
        if len(ind):
            k=ind[0];target=(max(-7,min(7,x[k]+1.5)),min(59,y[k]+5+j*2));ax.annotate(g,(x[k],y[k]),xytext=target,fontsize=5.5,arrowprops={'arrowstyle':'-','lw':.4,'color':'#555555'})
    ax.set(xlim=(-10,10),ylim=(0,64),xlabel='log2 fold change',ylabel='-log10 adjusted P (capped at 60)');clean(ax)
    ax.text(.03,.97,f'Down: {len(D["down"][c])}',transform=ax.transAxes,va='top',color=BLUE,fontsize=6)
    ax.text(.97,.97,f'Up: {len(D["up"][c])}',transform=ax.transAxes,va='top',ha='right',color=KD,fontsize=6)

# Figure 5: seven panels. Enrichment follows actual shared-DEG results, not the template's anticipated pathway list.
f=fig(210,270)
panel(f,.03,.95,'A','RNA-seq design');panel(f,.34,.95,'B','HT29');panel(f,.68,.95,'C','SW620')
for j,(a,b) in enumerate([('HT29 + SW620','shNC / shTCF7L2'),('RNA sequencing','3 samples per group / cell line'),('Differential expression','padj ≤ 0.05; |log2FC| ≥ 1'),('Shared programs','99 up / 17 down')]):
    y=.897-j*.046;text(f,.045,y,a,fontsize=7,fontweight='bold');text(f,.045,y-.016,b,fontsize=6)
    if j<3:text(f,.265,y-.029,'↓',fontsize=10,color=NC)
volcano(f.add_axes([.395,.735,.235,.18]),'HT29');volcano(f.add_axes([.735,.735,.235,.18]),'SW620')
panel(f,.03,.66,'D','Shared upregulated DEGs');panel(f,.37,.66,'E','Shared upregulated GO enrichment')
ax=f.add_axes([.03,.475,.30,.145]);ax.set_aspect('equal');ax.axis('off')
ax.add_patch(Circle((.39,.48),.30,color=NC,alpha=.23));ax.add_patch(Circle((.72,.48),.30,color=KD,alpha=.23));ax.set(xlim=(0,1.1),ylim=(0,1))
for x,s in [(.25,'276'),(.555,'99'),(.86,'1158')]:ax.text(x,.48,s,ha='center',va='center',fontsize=11,fontweight='bold')
ax.text(.24,.88,'HT29',ha='center',fontsize=8);ax.text(.86,.88,'SW620',ha='center',fontsize=8)
top=D['enrich']['up'].query('Count>0').head(8);dotenrich(f,[.735,.465,.205,.16],top,6)
text(f,.39,.434,'Black outline: FDR < 0.05; gray outline: FDR ≥ 0.05. Dot size reflects gene count.',fontsize=6)
panel(f,.03,.39,'F','Selected immune-associated genes');panel(f,.60,.39,'G','HT29 qPCR validation')
selected=['HLA-B','HLA-C','HLA-E','HLA-F','B2M','TAP1','IRF1','IFIT2','IFIT3','ISG20','GBP1','OASL','CXCL10','CXCL11','CCL5','CX3CL1','IL18','IFNB1']
idbyname={str(v):g for g,v in D['expr']['HT29'].gene_name.items()};ids=[idbyname[g] for g in selected if g in idbyname and idbyname[g] in D['expr']['SW620'].index]
heatmap(f,[.12,.12,.37,.225],ids,export='Fig5F_selected_heatmap')
ax=f.add_axes([.65,.17,.31,.17]);genes=['CCL5','CXCL10','IFNB1','IL18','HLA-B','HLA-C','ULBP-2']
for j,gene in enumerate(genes):
    for gi,v in enumerate(qpcr[gene]):
        x=j+(gi-.5)*.30;ax.scatter(x+np.linspace(-.055,.055,3),v,s=12,c=COLORS[gi],edgecolor='white',lw=.25)
        ax.errorbar(x,v.mean(),yerr=v.std(ddof=1)/np.sqrt(3),fmt='_',ms=8,c=COLORS[gi],capsize=2,lw=.7)
ax.set(xticks=range(7),xticklabels=genes,ylabel='Source-normalized mRNA expression',ylim=(0,7));ax.tick_params(axis='x',rotation=45,labelsize=6);clean(ax)
text(f,.65,.105,'3 sample means / group; paired technical wells averaged.\nMean ± SEM; no original six-well P values transferred.',fontsize=6.5,linespacing=1.5)
text(f,.045,.066,'Heatmap includes pathway-associated candidates; it does not imply all genes are shared significant DEGs.\nCXCL10, CXCL11 and IFNB1 are significant in HT29 only under the stated RNA-seq criterion.',fontsize=6.5,linespacing=1.5)
legend(f,.74,.368);footer(f,'GO: one-sided hypergeometric test; BH correction across all tested BP terms. 89 of 99 shared upregulated genes are BP-annotated.')
addfig('Fig5',f)

# S10: source PCA coordinates and source correlation matrices, no recomputed normalization.
f=fig(210,255);legend(f)
for ci,c in enumerate(CELLS):
    x=.11+ci*.50;panel(f,.035+ci*.50,.93,chr(65+ci),c+' | PCA')
    ax=f.add_axes([x,.715,.33,.18]);df=D['pca'][c]
    for gi,g in enumerate(['shNC','shTCF7L2']):
        sub=df[df.group==g];ax.scatter(sub.pc1,sub.pc2,s=30,c=COLORS[gi])
        for _,r in sub.iterrows():
            label=r['sample'].replace('shTCF7L2_','KD').replace('shNC_','NC')
            offset={'NC1':(-24,-3),'NC2':(5,-10),'NC3':(5,7)}.get(label,(4,3))
            ax.annotate(label,(r.pc1,r.pc2),xytext=offset,textcoords='offset points',fontsize=6,arrowprops={'arrowstyle':'-','lw':.35,'color':NC} if label.startswith('NC') else None)
    ax.margins(.23);ax.set(xlabel='PC1 (source coordinates)',ylabel='PC2 (source coordinates)');clean(ax)
    panel(f,.035+ci*.50,.62,chr(67+ci),c+' | Sample correlation')
    ax=f.add_axes([x,.385,.33,.195]);co=D['corr'][c].set_index('sample').loc[SAMPLECOLS,SAMPLECOLS].to_numpy();im=ax.pcolormesh(co,vmin=.94,vmax=1,cmap='YlGnBu');ax.set(xticks=np.arange(6)+.5,xticklabels=['NC1','NC2','NC3','KD1','KD2','KD3'],yticks=np.arange(6)+.5,yticklabels=['NC1','NC2','NC3','KD1','KD2','KD3'],ylim=(6,0));ax.tick_params(length=0,labelsize=6)
    for i in range(6):
        for j in range(6):ax.text(j+.5,i+.5,f'{co[i,j]:.3f}',ha='center',va='center',fontsize=5.5,color='white' if co[i,j]>.975 else 'black')
    cb=f.colorbar(im,ax=ax,fraction=.045,pad=.03);cb.solids.set_rasterized(False);cb.ax.tick_params(labelsize=6);cb.set_label('Correlation',fontsize=6)
    panel(f,.035+ci*.50,.29,chr(69+ci),c+' | Differential gene counts')
    ax=f.add_axes([x,.075,.33,.175]);counts=[len(D['up'][c]),len(D['down'][c])];ax.bar([0,1],counts,color=[KD,BLUE],width=.5)
    for j,n in enumerate(counts):ax.text(j,n+max(counts)*.03,str(n),ha='center',fontsize=8)
    ax.set(xticks=[0,1],xticklabels=['Upregulated','Downregulated'],ylabel='Number of genes',ylim=(0,max(counts)*1.2));clean(ax)
footer(f,'DESeq2 adjusted P ≤ 0.05 and |log2 fold change| ≥ 1. Six samples per cell line. PCA and correlations retain the delivered analysis.')
addfig('FigS10',f)

# S11: full result tables accompany a readable selection of pathways in the PDF.
f=fig(210,245)
for ci,c in enumerate(CELLS):
    y=.94-ci*.46;panel(f,.035,y,chr(65+ci),c+' | GO enrichment of all DEGs')
    df=D['go_original'][c];df=df[df.Category=='BP'].sort_values('padj').head(12).copy();df['GeneRatio']=df.GeneRatio.map(lambda s:float(s.split('/')[0])/float(s.split('/')[1]))
    dotenrich(f,[.60,y-.365,.33,.32],df,7)
footer(f,'Top 12 BP terms per cell line by original adjusted P value. Complete delivered tables are included in data/.')
addfig('FigS11',f)
f=fig(210,245)
for ci,c in enumerate(CELLS):
    y=.94-ci*.46;panel(f,.035,y,chr(67+ci),c+' | GSEA, non-apoptosis pathways')
    df=D['gsea'][c].copy();df['Description']=df.GENESET.str.split('#').str[-1].str.replace(r'\(G<LOCAL_PATH_REMOVED>','',regex=True).str.replace('_',' ').str.lower();df=df[~df.Description.str.contains('apopt')].sort_values(['FDR','NP']).head(12).iloc[::-1]
    ax=f.add_axes([.63,y-.365,.28,.32]);vals=df.NES.to_numpy();yy=np.arange(len(df));ax.barh(yy,vals,color=[KD if v>=0 else BLUE for v in vals],height=.6)
    ax.set(yticks=yy,yticklabels=['\n'.join(textwrap.wrap(s,52)) for s in df.Description],xlabel='Normalized enrichment score (NES)');ax.tick_params(axis='y',labelsize=6.5,length=0);clean(ax)
    for j,(_,r) in enumerate(df.iterrows()):ax.text(1.03,j,f'q={r.FDR:.3f}',transform=ax.get_yaxis_transform(),fontsize=6,va='center')
    ax.axvline(0,color='#666666',lw=.5)
footer(f,'Top 12 non-apoptosis gene sets by delivered FDR, then nominal P. Exact source NES / FDR retained; no new GSEA is inferred.')
addfig('FigS11',f)
f=fig(210,245)
for ci,direction in enumerate(['up','down']):
    y=.94-ci*.46;panel(f,.035,y,chr(69+ci),'Shared '+direction+'regulated DEGs | Recomputed GO')
    df=D['enrich'][direction].query('Count>0').head(12);dotenrich(f,[.61,y-.365,.32,.32],df,7)
    text(f,.045,y-.395,f'{len(D["common_"+direction])} shared genes; {int(df.QueryAnnotated.iloc[0])} BP-annotated; '+('3 terms pass FDR < 0.05.' if direction=='up' else 'No terms pass FDR < 0.05.'),fontsize=7)
footer(f,'Universe: 13,209 BP-annotated genes with nonmissing padj in both models; term size 10-500. Hypergeometric test + BH across all eligible terms.')
addfig('FigS11',f)

# S12: all requested module members quantified in BOTH matrices, paginated rather than silently omitted.
f=fig(210,280)
for ci,c in enumerate(CELLS):
    panel(f,.04+ci*.50,.95,chr(65+ci),c+' | Top 50 DEGs')
    d=D['deg'][c];sub=d[(d.padj<=.05)&(d.log2FoldChange.abs()>=1)].sort_values(['padj','pvalue']).head(50);ids=list(sub.index)
    heatmap(f,[.16+ci*.48,.105,.275,.79],ids,cells=[c],labelsize=6,export=f'S12_{c}_top50')
footer(f,'Top 50 by adjusted P, then nominal P among threshold-defined DEGs. log2(FPKM + 1), row z score within each cell line; colors clipped at ±2.')
addfig('FigS12',f)
for letter,title,gids in [('C','Antigen processing and presentation',['GO:0019882']),('D','Cytokine / chemokine activity',['GO:0005125','GO:0008009'])]:
    members=set().union(*(D['maps'][g] for g in gids));ids=sorted(members&set(D['expr']['HT29'].index)&set(D['expr']['SW620'].index),key=lambda g:str(D['expr']['HT29'].loc[g,'gene_name']))
    assert len(ids)==len(set(ids));perpage=80
    for start in range(0,len(ids),perpage):
        subset=ids[start:start+perpage];f=fig(210,280);panel(f,.04,.95,letter,title+f' | {start+1}-{start+len(subset)} / {len(ids)}')
        for ci in range(2):
            chunk=subset[ci*40:(ci+1)*40]
            if chunk:
                h=.77*len(chunk)/40
                heatmap(f,[.145+ci*.49,.88-h,.32,h],chunk,labelsize=5.8,export=f'S12_{letter}_{start+ci*40+1}',colorbar=False)
        text(f,.04,.05,'Source GO annotation: '+', '.join(gids)+'. All members quantified in both matrices are shown, including nonsignificant genes.',fontsize=6.5)
        footer(f,'Within-cell-line row z scores of log2(FPKM + 1); blue = -2, white = 0, red = +2. Zero-variance rows are white; colors clipped at ±2.')
        addfig('FigS12',f)

# Export PDF with editable text / paths and native SVG for each page; photos remain raster.
for name,figures in FIGURES.items():
    with PdfPages(OUT/(name+'.pdf')) as pdf:
        for i,f in enumerate(figures):
            pdf.savefig(f);f.savefig(OUT/(name+(f'_p{i+1}' if len(figures)>1 else '')+'.svg'));plt.close(f)
    d=fitz.open(OUT/(name+'.pdf'))
    for i,p in enumerate(d):p.get_pixmap(matrix=fitz.Matrix(1.25,1.25)).save(str(TMP/(name+f'_review_p{i+1}.png')))
    print('EXPORTED',name,len(d),flush=True)
pd.DataFrame(STATS).to_csv(DATA/'new_endpoint_statistics.csv',index=False)
PROVENANCE['GO_reanalysis']={'universe':len(D['universe']),'size_filter':[10,500],'test':'hypergeometric upper-tail','correction':'BH across all eligible BP terms, including zero overlap','up_n':len(D['common_up']),'down_n':len(D['common_down'])}
PROVENANCE['representative_TUNEL']={g:{'tissue_id':v[0],'field':v[1],'path':source(v[2])} for g,v in representatives.items()}
PROVENANCE['animal_method']={'source':'论文.pdf','PDF_pages':[36,50,57],'note':'strain/sex supported; inoculum dose internally inconsistent and omitted'}
(OUT/'source_map.json').write_text(json.dumps(PROVENANCE,ensure_ascii=False,indent=2),encoding='utf-8')

