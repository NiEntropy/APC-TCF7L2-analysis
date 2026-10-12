from pathlib import Path
import sys
ROOT=Path.cwd();sys.path.insert(0,str(ROOT/'tmp/pdfs/python_libs'))
import json,hashlib,shutil,zipfile,re
import xml.etree.ElementTree as ET
import numpy as np
import pandas as pd
import openpyxl
from scipy import stats
from scipy.stats import ttest_ind
from PIL import Image
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import pymupdf as fitz

SRC=ROOT/'2026 崔珣 合成致死课题原始数据'
OUT=ROOT/'output/pdf/Fig4_ANOVA_Sidak_20261008';OUT.mkdir(parents=True,exist_ok=True)
DATA=OUT/'data';DATA.mkdir(exist_ok=True)
NC='#777D83';KD='#8A3048';COLORS=[NC,KD];STATS=[]
plt.rcParams.update({'font.family':'Arial','font.size':7,'pdf.fonttype':42,'svg.fonttype':'none','axes.linewidth':.6,'xtick.major.size':2,'ytick.major.size':2,'axes.labelpad':3,'savefig.facecolor':'white'})
animalpath=next(SRC.rglob('HT29接瘤数据第二批原始数据.xlsx'))
s=openpyxl.load_workbook(animalpath,data_only=True).active
days=np.array([4,7,9,11,14,16,18,20,22]);bodydays=np.r_[0,days];animal={}
for gi,g in enumerate(['shNC','shTCF7L2']):
    rr=range(3+8*gi,11+8*gi)
    animal[g]={'id':[s.cell(r,2).value for r in rr], 'volume':np.array([[s.cell(r,c).value for c in range(7,40,4)] for r in rr],float),'weight':np.array([s.cell(r,40).value for r in rr],float)}
v=np.vstack([animal[g]['volume'] for g in animal]);assert v.shape==(16,9) and np.isfinite(v).all()
pd.DataFrame([[g,id,int(day),float(value)] for g,a in animal.items() for id,vs in zip(a['id'],a['volume']) for day,value in zip(days,vs)],columns=['group','animal_ID_as_recorded','day','volume_mm3']).to_csv(DATA/'Fig4C_individual_volumes.csv',index=False)

# Reproduce the supplied ordinary two-way ANOVA with the Prism Day0 zero row.
# Days4-22 use the eight-animal Excel records whose group means match the screenshot.
for g in animal:animal[g]['volume']=np.c_[np.zeros(8),animal[g]['volume']]
days=np.r_[0,days];v=np.vstack([animal[g]['volume'] for g in animal])
a=np.stack([animal[g]['volume'] for g in animal]);m=a.mean(axis=1);grand=a.mean();gm=m.mean(axis=1);tm=m.mean(axis=0)
SSg=8*10*np.sum((gm-grand)**2);SSt=8*2*np.sum((tm-grand)**2);SSgt=8*np.sum((m-gm[:,None]-tm+grand)**2);SSe=np.sum((a-m[:,None,:])**2);mse=SSe/140
assert np.isclose(np.sum((a-grand)**2),SSg+SSt+SSgt+SSe)
print('SSE check',SSe);assert abs(SSe-3556535)<10
ANOVA=[]
for name,ss,df in [('Interaction',SSgt,9),('Time',SSt,9),('Group',SSg,1)]:
    F=ss/df/mse;ANOVA.append(dict(effect=name,SS=ss,df=df,MS=ss/df,F=F,df_error=140,pvalue=float(stats.f.sf(F,df,140))))
ANOVA.append(dict(effect='Residual',SS=SSe,df=140,MS=mse))
pd.DataFrame(ANOVA).to_csv(DATA/'Fig4C_ordinary_two_way_ANOVA.csv',index=False)
se=np.sqrt(mse*(1/8+1/8));diff=m[0]-m[1];ts=diff/se;raw=2*stats.t.sf(abs(ts),140)
with np.errstate(divide='ignore'):adj=-np.expm1(10*np.log1p(-raw))
critical=stats.t.ppf(1-(1-.95**.1)/2,140);margin=critical*se
POST=pd.DataFrame({'day':days,'shNC_mean':m[0],'shTCF7L2_mean':m[1],'mean_difference_NC_minus_KD':diff,'pooled_SE':se,'t':ts,'df':140,'p_raw':raw,'p_Sidak_10_days':adj,'simultaneous_CI95_low':diff-margin,'simultaneous_CI95_high':diff+margin})
POST['star']=np.where(adj<.0001,'****',np.where(adj<.001,'***',np.where(adj<.01,'**',np.where(adj<.05,'*','ns'))))
assert round(adj[7],4)==.0176 and round(adj[8],4)==.0002 and adj[9]<.0001
assert list(POST.star)==['ns']*7+['*','***','****']
assert round(se,2)==79.69
POST.to_csv(DATA/'Fig4C_Sidak_comparisons.csv',index=False)
pd.DataFrame([[g,id,int(day),float(value),'Prism zero baseline' if day==0 else 'Excel measurement record'] for g,x in animal.items() for id,vs in zip(x['id'],x['volume']) for day,value in zip(days,vs)],columns=['group','animal_ID_as_recorded','day','volume_mm3','source']).to_csv(DATA/'Fig4C_individual_volumes.csv',index=False)

# D: use the Prism table corresponding to the user's supplied results, not edited Excel data.
prism=next(SRC.rglob('HT29 第二批.pzfx'));ns={'p':'htt<LOCAL_PATH_REMOVED>'}
t=next(t for t in ET.parse(prism).getroot().findall('p:Table',ns) if t.findtext('p:Title',namespaces=ns)=='Tumor weight')
weights=[np.array([float(d.text) for d in y.findall('.//p:d',ns)]) for y in t.findall('p:YColumn',ns)]
assert [len(a) for a in weights]==[8,8]
testD=ttest_ind(*weights,equal_var=True);diff=weights[1].mean()-weights[0].mean();sp2=sum(7*np.var(a,ddof=1) for a in weights)/14;sediff=np.sqrt(sp2/4);ci=diff+np.array([-1,1])*stats.t.ppf(.975,14)*sediff
fvar=np.var(weights[0],ddof=1)/np.var(weights[1],ddof=1);fp=2*min(stats.f.cdf(fvar,7,7),stats.f.sf(fvar,7,7))
RESULTD={'method':'two-sided unpaired Student t test, equal variances','n_shNC':8,'n_shTCF7L2':8,'mean_shNC':weights[0].mean(),'mean_shTCF7L2':weights[1].mean(),'t':float(testD.statistic),'df':14,'p':float(testD.pvalue),'difference_KD_minus_NC':diff,'difference_SEM':sediff,'CI95_low':ci[0],'CI95_high':ci[1],'variance_F':fvar,'variance_F_p':fp}
assert round(RESULTD['p'],4)==.0482 and round(RESULTD['t'],3)==2.165
assert round(RESULTD['mean_shNC'],4)==.8830 and round(RESULTD['mean_shTCF7L2'],4)==.4704
assert round(fp,4)==.1562
pd.DataFrame([RESULTD]).to_csv(DATA/'Fig4D_Student_t_test.csv',index=False)
pd.DataFrame({'Prism_record_index':range(1,9),'shNC_g':weights[0],'shTCF7L2_g':weights[1]}).to_csv(DATA/'Fig4D_Prism_weights.csv',index=False)
excelweights=[a['weight'].copy() for a in animal.values()]
pd.DataFrame([{'dataset':label,'test':method,'P':ttest_ind(*arrays,equal_var=eq).pvalue} for label,arrays in [('Excel',excelweights),('Prism',weights)] for method,eq in [('Student (equal variance)',True),('Welch',False)]]).to_csv(DATA/'Fig4D_method_and_source_comparison.csv',index=False)
for g,w in zip(animal,weights):animal[g]['weight']=w

# Preserve the existing F panel and its Welch test.
tpath=next(SRC.rglob('TUNEL分析.xlsx'));s=openpyxl.load_workbook(tpath,data_only=True).active;tunel={}
for g,starts in [('shNC',range(2,26,3)),('shTCF7L2',range(28,52,3))]:
    ids=[s.cell(r,1).value for r in starts];areas=np.array([[s.cell(r+j,7).value for j in range(3)] for r in starts],float)
    tunel[g]={'id':ids,'fields':areas,'area':areas.mean(1)}

# Reuse the established layout functions, embedded into this self-contained revision builder.

FIGURES={};HEATMAP_ROWS=[]
def fig(w=210,h=230):return plt.figure(figsize=(w/25.4,h/25.4))
def text(f,x,y,s,**kw):return f.text(x,y,s,**kw)
def panel(f,x,y,letter,title):
    text(f,x,y,letter,fontsize=11,fontweight='bold');text(f,x+.033,y,title,fontsize=8,fontweight='bold')
def legend(f,x=.76,y=.978):
    text(f,x,y,'● shNC',color=NC,fontsize=7);text(f,x+.105,y,'● shTCF7L2',color=KD,fontsize=7)
def clean(ax):ax.spines[['top','right']].set_visible(False);ax.tick_params(labelsize=7)
def footer(f,s):text(f,.04,.018,s,fontsize=6,color='#454545')
def addfig(name,f):FIGURES.setdefault(name,[]).append(f)
def stars(p):return '****' if p<.0001 else '***' if p<.001 else '**' if p<.01 else '*' if p<.05 else 'ns'
def dots(ax,arrays,ylabel,labels=('shNC','shTCF7L2'),test=None):
    vmax=max(np.max(v) for v in arrays)
    for i,v in enumerate(arrays):
        v=np.asarray(v);ax.scatter(i+np.linspace(-.10,.10,len(v)),v,s=18,c=COLORS[i],edgecolor='white',lw=.35,zorder=3)
        ax.errorbar(i,np.mean(v),yerr=np.std(v,ddof=1)/np.sqrt(len(v)),fmt='_',color='black',capsize=4,ms=17,lw=.8,zorder=4)
    ax.set(xticks=[0,1],xticklabels=labels,ylabel=ylabel,xlim=(-.45,1.45),ylim=(0,vmax*1.28));clean(ax)
    if test:
        equal_var=test.startswith('Fig4D');p=float(ttest_ind(*arrays,equal_var=equal_var).pvalue);STATS.append(dict(metric=test,test='Student' if equal_var else 'Welch',pvalue=p,n1=len(arrays[0]),n2=len(arrays[1])))
        y=vmax*1.10;ax.plot([0,0,1,1],[y,y+vmax*.03,y+vmax*.03,y],c='black',lw=.6);ax.text(.5,y+vmax*.05,stars(p) if p<.05 else f'P = {p:.3f}',ha='center',fontsize=8)
def imageax(f,box,p,crop=None,title=None):
    im=Image.open(p);im=im.crop(crop) if crop else im
    ax=f.add_axes(box);ax.imshow(im);ax.axis('off')
    if title:ax.set_title(title,fontsize=7,pad=4)
    return ax
def growth(ax,key='volume',individual=False):
    times=days if key=='volume' else bodydays
    for gi,(g,a) in enumerate(animal.items()):
        vs=a[key]
        if individual:
            for id,v in zip(a['id'],vs):ax.plot(times,v,lw=.7,color=COLORS[gi],alpha=.65)
        else:ax.errorbar(times,vs.mean(axis=0),yerr=vs.std(axis=0,ddof=1)/np.sqrt(8),color=COLORS[gi],fmt='o-',ms=3,capsize=2,lw=1)
    ax.set(xlabel='Days after implantation',ylabel='Tumor volume (mm³)' if key=='volume' else 'Body weight (g)',xticks=[0,7,14,22],xlim=(0,23));clean(ax)
    if key=='volume':ax.set_ylim(bottom=0)
def select_tunel(g,id,field=1):
    root=SRC/'Figure 4/Figure 4-4/2 皮下瘤组织 免疫荧光/2 IF拍摄原始照片'/g
    candidates=[p for p in root.glob('*.jpg') if re.search(rf'\s{id}肿-Merge_10\.00X-{field}\.jpg$',p.name,re.I)]
    if len(candidates)!=1:raise ValueError((g,id,field,candidates))
    return candidates[0]
photo=SRC/'Figure 3/Figure 3-3 动物实验/1 解剖结果/解剖肿瘤.jpg'

# Fig4. No Ki-67 quantification is mixed in because its sample labels are unresolved.
f=fig(180,195);legend(f,.71,.98)
panel(f,.035,.94,'A','HT29 xenograft design')
steps=[('HT29 cells','shNC / shTCF7L2'),('Subcutaneous implantation','Male BALB/c nude mice'),('Growth monitoring','8 mice / group'),('Endpoint analysis','Day 23')]
for i,(a,b) in enumerate(steps):
    y=.872-i*.046;text(f,.07,y,a,fontsize=7,fontweight='bold');text(f,.07,y-.018,b,fontsize=6.5,color='#555555')
    if i<3:text(f,.39,y-.03,'↓',fontsize=11,color=NC)
panel(f,.49,.94,'B','Excised tumors')
imageax(f,[.52,.70,.44,.22],photo,crop=(75,510,1040,930))
text(f,.52,.69,'Upper: shNC | Lower: shTCF7L2 (8 tumors each)',fontsize=6)
panel(f,.035,.64,'C','Tumor growth');panel(f,.54,.64,'D','Endpoint tumor weight')
axC=f.add_axes([.12,.40,.34,.205]);growth(axC)
axC.set_ylim(0,1450)
axC.text(.035,.95,'Ordinary two-way ANOVA\nGroup x time: P < 0.0001',transform=axC.transAxes,va='top',fontsize=6.5,linespacing=1.4)
for j,row in POST.iterrows():
    if row['p_Sidak_10_days']<.05:
        highest=max(a['volume'][:,j].mean()+a['volume'][:,j].std(ddof=1)/np.sqrt(8) for a in animal.values())
        axC.text(row['day'],highest+60,row['star'],ha='center',va='bottom',fontsize=10)

dots(f.add_axes([.64,.40,.29,.205]),[a['weight'] for a in animal.values()],'Tumor weight (g)',test='Fig4D tumor weight, day23')
panel(f,.035,.32,'E','TUNEL / DAPI');panel(f,.69,.32,'F','TUNEL-positive area')
representatives={}
for gi,(g,a) in enumerate(tunel.items()):
    order=np.argsort(a['area']);ix=order[(len(order)-1)//2];id=a['id'][ix];field=int(np.argmin(abs(a['fields'][ix]-a['area'][ix])))+1;p=select_tunel(g,id,field);representatives[g]=(id,field,p)
    imageax(f,[.045+gi*.31,.11,.295,.17],p,title=g+f' | tissue {id}')
dots(f.add_axes([.77,.11,.18,.16]),[a['area'] for a in tunel.values()],'Positive area (%)',test='Fig4F TUNEL area, sample means')
text(f,.045,.069,'C: ordinary two-way ANOVA with Sidak-adjusted row comparisons (10 comparisons, one family).',fontsize=5.8)
text(f,.045,.049,'D: two-sided unpaired Student t test (Prism data), P = 0.0482; F: Welch test. Mean ± SEM.',fontsize=5.8)
footer(f,'C: Day18 *, Day20 ***, Day22 ****. This reproduction uses the supplied non-repeated-measures model.')
addfig('Fig4',f)

f.savefig(OUT/'Fig4.pdf');f.savefig(OUT/'Fig4.svg');plt.close(f)
doc=fitz.open(OUT/'Fig4.pdf');page=doc[0]
assert all(b[0]>=-1 and b[1]>=-1 and b[2]<=page.rect.width+1 and b[3]<=page.rect.height+1 for b in page.get_text('blocks'))
page.get_pixmap(matrix=fitz.Matrix(1.8,1.8)).save(str(OUT/'Fig4_preview.png'))
pd.DataFrame(STATS).to_csv(DATA/'Fig4_endpoint_tests.csv',index=False)
sources={'C':str(animalpath.relative_to(ROOT)),'D':str(prism.relative_to(ROOT))+' / Tumor weight','F':str(tpath.relative_to(ROOT)),'D_screenshot':'codex-clipboard-f117e51c-1b23-4a21-add7-72621ff2e29f.png','C_method':'ordinary two-way ANOVA; Sidak 10 comparisons','D_verification':RESULTD}
(OUT/'source_and_statistics.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2),encoding='utf-8')
print('C ordinary ANOVA:',ANOVA)
print('C posthoc:',POST.to_string(index=False))
print('D screenshot replicated:',RESULTD)
print('Figure saved:',OUT/'Fig4.pdf')



