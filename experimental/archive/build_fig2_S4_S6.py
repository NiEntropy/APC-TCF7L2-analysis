from pathlib import Path
import xml.etree.ElementTree as ET
import json, csv, sys
sys.path.insert(0,str(Path.cwd()/'tmp/pdfs/python_libs'))
import xlrd
import numpy as np
import openpyxl
from PIL import Image
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch
import fitz

ROOT=Path.cwd()
SRC=ROOT/'2026 崔珣 合成致死课题原始数据'/'Figure 3'
OUT=ROOT/'output'/'pdf'/'Fig2_S4_S6'
OUT.mkdir(parents=True,exist_ok=True)
NS={'p':'htt<LOCAL_PATH_REMOVED>'}
ORDER=['HT29','SW620','LOVO','LS174T','RKO','NCM460']
DIRS={c:next(p for p in SRC.glob('*/*') if p.is_dir() and p.name.lower().startswith(c.lower()+'-')) for c in ORDER}
plt.rcParams.update({'font.family':'Arial','font.size':7,'pdf.fonttype':42,'ps.fonttype':42,'svg.fonttype':'none','axes.linewidth':.65,'xtick.major.width':.65,'ytick.major.width':.65,'xtick.major.size':2.5,'ytick.major.size':2.5,'axes.labelpad':2,'savefig.facecolor':'white'})
RED='#CC514B'; BLUE='#407AA6'; ORANGE='#D79943'; GRAY='#737980'; INK='#20252B'
palette={'NC':GRAY,'#1':RED,'#4':ORANGE,'#6':RED}
sources={}; growth={}; wb={}; crops={}; qpcr={}; colony={}; colony_images={}
SIGNIF={
 'WB':{'HT29':{'#1':'****'},'SW620':{'#4':'ns','#6':'**'},'LOVO':{'#4':'**','#6':'**'},'LS174T':{'#1':'****'},'RKO':{'#6':'**'},'NCM460':{'#6':'**'}},
 'qPCR':{'HT29':{'#1':'***'},'SW620':{'#4':'****','#6':'****'},'LOVO':{'#4':'****','#6':'****'},'LS174T':{'#1':'***'},'RKO':{'#6':'****'},'NCM460':{'#6':'****'}},
 'SRB':{'HT29':{'#1':'****'},'SW620':{'#6':'****'},'LOVO':{'#4':'****','#6':'****'},'LS174T':{'#1':'ns'},'RKO':{'#6':'ns'},'NCM460':{'#6':'ns'}},
 'colony':{'HT29':{'#1':'****'},'SW620':{'#6':'****'},'LOVO':{'#4':'***','#6':'***'},'LS174T':{'#1':'ns'},'RKO':{'#6':'ns'},'NCM460':{'#6':'ns'}}}
def xml_data(p,cell=None):
 r=ET.parse(p).getroot(); tables=r.findall('p:Table',NS)
 t=next(t for t in tables if t.findtext('p:Title',namespaces=NS).strip().upper()==cell.upper()) if cell else tables[0]
 groups={}
 for y in t.findall('p:YColumn',NS):
  name=''.join(y.find('p:Title',NS).itertext()).strip(); label='NC' if name=='shNC' else '#'+name.split('-')[-1]
  groups[label]=np.array([[float(d.text) for d in sub.findall('p:d',NS)] for sub in y.findall('p:Subcolumn',NS)]).T
 return groups
for c in ORDER:
 p=next(DIRS[c].rglob('*SRB.pzfx'))
 if c!='RKO': growth[c]=xml_data(p)
 else:
  p=next(DIRS[c].rglob('*SRB*.xlsx')); book=openpyxl.load_workbook(p,data_only=True)
  growth[c]={label:np.array([[book[f'{h}h'].cell(r,col).value for r in range(12,18)] for h in [12,24,48,72,96]]) for label,col in [('NC',8),('#6',9)]}
 sources[c]={'growth':str(p.relative_to(ROOT))}
for c in ORDER:
 p=next(p for p in DIRS[c].rglob('*蛋白灰度分析.xlsx') if not p.name.startswith('~')); sheet=openpyxl.load_workbook(p,data_only=True).active
 # Recalculate target/loading-control ratios from the original columns,
 # rather than relying on the mislabeled LOVO Prism table.
 rows=list(sheet.iter_rows(values_only=True)); groups={}; active=None
 for row in rows[1:]:
  vals=[v for v in row if v is not None]
  if not vals:continue
  if isinstance(vals[0],str):
   active='NC' if vals[0].strip()=='shNC' else '#'+vals[0].strip().split('-')[-1]
   groups[active]=[]; vals=vals[1:]
  if active and len(vals)>=5: groups[active].append(float(vals[3])/float(vals[1]))
 den=np.mean(groups['NC']); wb[c]={g:np.array(v)/den for g,v in groups.items()}
 sources[c]['densitometry']=str(p.relative_to(ROOT))
boxes={('HT29','TCF7L2'):(386,526,963,568),('HT29','GAPDH'):(437,491,1028,536),('SW620','TCF7L2'):(451,577,887,620),('SW620','GAPDH'):(429,566,876,611),('LOVO','TCF7L2'):(480,557,910,620),('LOVO','GAPDH'):(430,565,878,613),('LS174T','TCF7L2'):(380,508,950,562),('LS174T','GAPDH'):(359,507,948,554),('NCM460','TCF7L2'):(399,471,978,510),('NCM460','GAPDH'):(471,452,1073,499),('RKO','TCF7L2'):(601,413,883,454),('RKO','GAPDH'):(805,466,1094,516)}
for (c,protein),box in boxes.items():
 p=next(p for p in DIRS[c].rglob('*signal.jpg') if p.parent.name==protein)
 crops[c,protein]=Image.open(p).crop(box)
 sources[c][protein]={'file':str(p.relative_to(ROOT)),'crop_xyxy':box,'adjustment':'none; rectangular crop only; no lane splicing'}

for c in ORDER:
 if c!='RKO':
  p=next(p for p in DIRS[c].rglob('*.pzfx') if 'mRNA' in str(p));qpcr[c]={g:a.flatten() for g,a in xml_data(p,c).items()}
 else:
  p=next(p for p in DIRS[c].rglob('*.xls') if not p.name.startswith('~'));s=xlrd.open_workbook(p).sheet_by_name('Sheet1')
  qpcr[c]={'NC':np.array([s.cell_value(i,11) for i in range(13,19)]),'#6':np.array([s.cell_value(i,11) for i in range(19,25)])}
 sources[c]['qPCR']=str(p.relative_to(ROOT))
 if c=='RKO':sources[c]['qPCR_range']='Sheet1!L14:L19 (NC-puro), L20:L25 (sh6); stored normalized results'
 p=next(p for p in DIRS[c].rglob('克隆统计.xlsx') if not p.name.startswith('~'))
 sh=openpyxl.load_workbook(p,data_only=True).active;gg={};active=None
 for row in sh.iter_rows(values_only=True):
  vals=[v for v in row if v is not None]
  if not vals:continue
  if isinstance(vals[0],str):
   active='NC' if vals[0].strip()=='shNC' else '#'+vals[0].strip().split('-')[-1];gg[active]=[];vals=vals[1:]
  if active and vals and isinstance(vals[0],(int,float)):gg[active].append(float(vals[0]))
 colony[c]={g:np.array(v) for g,v in gg.items()};sources[c]['colony_counts']=str(p.relative_to(ROOT))
 ims={}
 for g in colony[c]:
  pp=sorted(p for p in DIRS[c].rglob('*.png') if '克隆形成' in str(p) and (('shNC' in p.name) if g=='NC' else (('TCFsh1' in p.name) if c=='HT29' else (('TCF7L2'+g) in p.name or ('TCF7L2-'+g[1:]) in p.name))))
  assert len(pp)==3,(c,g,pp)
  ims[g]=pp
 colony_images[c]=ims;sources[c]['colony_images']={g:[str(p.relative_to(ROOT)) for p in pp] for g,pp in ims.items()}

W,H=180,125
fig=plt.figure(figsize=(W/25.4,H/25.4))
def txt(x,y,s,size=7,weight='normal',color=INK,ha='left',**kw):
 return fig.text(x/W,y/H,s,fontsize=size,fontweight=weight,color=color,ha=ha,va='center',**kw)
def axmm(x,y,w,h):return fig.add_axes([x/W,y/H,w/W,h/H])
def clean(ax):
 ax.spines[['top','right']].set_visible(False);ax.tick_params(labelsize=6)
def title(letter,x,y,text):
 txt(x,y,letter,10,'bold');txt(x+6,y,text,8,'bold')
def bracket(ax,x1,x2,y,text,dy=.06,size=7):
 ax.plot([x1,x1,x2,x2],[y-dy,y,y,y-dy],color=INK,lw=.7,clip_on=False)
 ax.text((x1+x2)/2,y+dy*.14,text,ha='center',va='bottom',fontsize=size,color=INK)

title('A',5,120,'Experimental design')
txt(11,112,'Cell line',6.5,'bold');txt(34,112,'Model group',6.5,'bold');txt(62,112,'Growth shRNA',5.8,'bold')
for i,c in enumerate(ORDER):
 y=107-i*3.6; col=RED if i<3 else (ORANGE if c=='LS174T' else BLUE)
 txt(11,y,c,6.6,'bold',col)
 txt(34,y,'APC-mutant CRC' if i<3 else ('Comparator CRC' if c!='NCM460' else 'Normal intestinal'),6.3)
 txt(64,y,'#1' if c in ['HT29','LS174T'] else ('#4, #6' if c=='LOVO' else '#6'),6.3)
txt(11,81.8,'shNC / shTCF7L2',6.5)
fig.add_artist(FancyArrowPatch((40/W,81.8/H),(47/W,81.8/H),transform=fig.transFigure,arrowstyle='->',mutation_scale=6,lw=.6,color=GRAY))
txt(49,81.8,'SRB + colony formation',6.3)

title('B',82,120,'TCF7L2 knockdown validation')
txt(89,106.5,'TCF7L2',5.3,ha='right');txt(89,101.8,'GAPDH',5.3,ha='right')
for i,c in enumerate(ORDER[:3]):
 x=91+i*29; width=25
 txt(x+width/2,113.5,c,7,'bold',ha='center')
 groups=list(wb[c]);n=len(groups)
 for j,g in enumerate(groups):txt(x+(j+.5)*width/n,110,g,5.3,ha='center')
 for protein,y in [('TCF7L2',104.6),('GAPDH',100)]:
  ax=axmm(x,y,width,3.8);ax.imshow(crops[c,protein],aspect='auto',interpolation='none');ax.axis('off')
 ax=axmm(x+2,85.5,width-3,11.5);clean(ax)
 for j,(g,vals) in enumerate(wb[c].items()):
  col=palette[g];ax.bar(j,vals.mean(),width=.57,color=col,alpha=.24,edgecolor=col,lw=.6)
  ax.errorbar(j,vals.mean(),yerr=vals.std(ddof=1),color=col,capsize=1.8,lw=.7)
  ax.scatter(j+np.linspace(-.14,.14,len(vals)),vals,s=6,color=col,zorder=4,lw=0)
 ax.set_xticks(range(n),groups,fontsize=5.5);ax.set_ylim(0,1.85);ax.set_yticks([0,1]);ax.tick_params(pad=1,labelsize=5.3)
 for j,g in enumerate(groups[1:],1):bracket(ax,0,j,1.36+(j-1)*.30,SIGNIF['WB'][c][g],dy=.07,size=5.6)
 if i==0:ax.set_ylabel('Relative protein',fontsize=5.5,labelpad=1)

times=np.array([12,24,48,72,96])
for i,c in enumerate(ORDER[:3]):
 x=[15,73,131][i];title('CDE'[i],x-10,75.5,c)
 ax=axmm(x,45,41,25);clean(ax)
 for g,a in growth[c].items():
  col=palette[g];ax.errorbar(times,a.mean(axis=1),yerr=a.std(axis=1,ddof=1),color=col,marker='o',markersize=2.5,lw=1,capsize=1.6,label='shNC' if g=='NC' else 'shTCF7L2 '+g+'  '+SIGNIF['SRB'][c][g])
 ax.set_xticks(times);ax.set_xlim(7,101);ax.set_ylim(0,1.85);ax.set_yticks([0,.5,1,1.5]);ax.set_xlabel('Time (h)',fontsize=6.5)
 if i==0:ax.set_ylabel('SRB absorbance (575 nm)',fontsize=6.5)
 ax.legend(frameon=False,fontsize=5.3,loc='upper left',handlelength=1.4,borderaxespad=.2,labelspacing=.3)

title('F',5,35,'Relative SRB signal at 96 h')
ax=axmm(20,10.5,152,19);clean(ax)
positions=[0,1,1.85,2.45,3.5,4.5,5.5]
conditions=[('HT29','#1'),('SW620','#6'),('LOVO','#4'),('LOVO','#6'),('LS174T','#1'),('RKO','#6'),('NCM460','#6')]
from scipy.stats import ttest_ind
endpoint_p=np.array([ttest_ind(growth[c][g][-1],growth[c]['NC'][-1],equal_var=False).pvalue for c,g in conditions])
endpoint_order=np.argsort(endpoint_p);endpoint_adjusted=np.empty(7)
endpoint_adjusted[endpoint_order]=np.minimum(1,np.maximum.accumulate(endpoint_p[endpoint_order]*(7-np.arange(7))))
endpoint_stars={k:('****' if p<.0001 else '***' if p<.001 else '**' if p<.01 else '*' if p<.05 else 'ns') for k,p in zip(conditions,endpoint_adjusted)}
export=[]
for x,(c,g) in zip(positions,conditions):
 vals=100*growth[c][g][-1]/growth[c]['NC'][-1].mean();col=RED if c in ORDER[:3] else (ORANGE if c=='LS174T' else BLUE)
 ax.bar(x,vals.mean(),width=.44,color=col,alpha=.23,edgecolor=col,lw=.7)
 ax.errorbar(x,vals.mean(),yerr=vals.std(ddof=1),color=col,capsize=2,lw=.8)
 ax.scatter(x+np.linspace(-.12,.12,len(vals)),vals,s=8,facecolor='white' if g=='#4' else col,edgecolor=col,lw=.6,zorder=4)
 ax.text(x,max(vals.max(),vals.mean()+vals.std(ddof=1))+9,endpoint_stars[c,g],ha='center',va='baseline',fontsize=7)
 for j,v in enumerate(vals):export.append([c,g,j+1,float(growth[c][g][-1,j]),float(growth[c]['NC'][-1].mean()),float(v)])
ax.axhline(100,color=GRAY,ls=(0,(3,2)),lw=.7,zorder=0)
ax.set_ylim(0,125);ax.set_yticks([0,50,100]);ax.set_ylabel('% of matched shNC',fontsize=6.5)
ax.set_xticks(positions,[c+'\n'+g for c,g in conditions],fontsize=6);ax.tick_params(axis='x',length=0,pad=2)
ax.set_xlim(-.5,6);ax.text(6,103,'shNC',fontsize=5.5,color=GRAY,ha='right')
pdf=OUT/'Fig2_editable.pdf'
txt(5,2.7,'F: two-sided Welch t-test; Holm correction (7 comparisons); technical-well analysis.',5,color=GRAY)
fig.savefig(pdf);fig.savefig(OUT/'Fig2_editable.svg');fig.savefig(OUT/'Fig2_preview.png',dpi=300)
with (OUT/'Fig2_source_data.csv').open('w',encoding='utf-8-sig',newline='') as f:
 writer=csv.writer(f);writer.writerow(['cell_line','shRNA','well','SRB_96h','matched_NC_mean_96h','percent_of_NC']);writer.writerows(export)
with (OUT/'Fig2_growth_data.csv').open('w',encoding='utf-8-sig',newline='') as f:
 writer=csv.writer(f);writer.writerow(['cell_line','shRNA','time_h','well','absorbance_575nm'])
 for c in ORDER:
  for g,a in growth[c].items():
   for t,row in zip(times,a):
    for j,v in enumerate(row):writer.writerow([c,g,int(t),j+1,float(v)])
(OUT/'Fig2_source_map.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2),encoding='utf-8')
doc=fitz.open(pdf);page=doc[0];page.get_pixmap(matrix=fitz.Matrix(3,3)).save(str(ROOT/'tmp/pdfs/Fig2_render.png'))
print('PDF',page.rect,'images',len(page.get_images()),'vector paths',len(page.get_drawings()),'text chars',len(page.get_text()))
assert len(page.get_images())==6
assert len(page.get_drawings())>50
assert all(c in page.get_text() for c in ORDER)
print('Created:',pdf)

def start(w,h):
 global fig,W,H
 W,H=w,h;fig=plt.figure(figsize=(w/25.4,h/25.4))
def save(name):
 fig.savefig(OUT/(name+'_editable.pdf'));fig.savefig(OUT/(name+'_editable.svg'));fig.savefig(OUT/(name+'_preview.png'),dpi=300)
 doc=fitz.open(OUT/(name+'_editable.pdf'));pg=doc[0];pg.get_pixmap(matrix=fitz.Matrix(2.5,2.5)).save(str(ROOT/'tmp/pdfs'/f'{name}_render.png'))
 assert len(pg.get_drawings())>20
 assert len(pg.get_text())>50
 print(name, 'size',pg.rect,'images',len(pg.get_images()),'paths',len(pg.get_drawings()))
def barplot(ax,data,c,assay,ylabel,ylim=None):
 clean(ax); groups=list(data)
 for j,(g,a) in enumerate(data.items()):
  col=palette[g];ax.bar(j,a.mean(),width=.58,color=col,alpha=.23,edgecolor=col,lw=.65)
  ax.errorbar(j,a.mean(),yerr=a.std(ddof=1),color=col,lw=.8,capsize=2)
  ax.scatter(j+np.linspace(-.13,.13,len(a)),a,s=10,color=col,lw=0,zorder=3)
 peak=max(max(a) for a in data.values());limit=ylim or (peak*1.55)
 ax.set_ylim(0,limit);ax.set_xlim(-.5,len(groups)-.5);ax.set_xticks(range(len(groups)),groups,fontsize=6)
 ax.set_ylabel(ylabel,fontsize=6,labelpad=2);ax.tick_params(labelsize=5.5,pad=2)
 for j,g in enumerate(groups[1:],1):bracket(ax,0,j,peak+limit*(.07+.12*(j-1)),SIGNIF[assay][c][g],dy=limit*.025,size=6.5)
def blots(c,x,top,width):
 if c=='LS174T':groups=[('#1',1/12),('NC',5/12),('#1',10/12)]
 else:groups=[(g,(j+.5)/len(wb[c])) for j,g in enumerate(wb[c])]
 for g,frac in groups:txt(x+frac*width,top+3,g,5.7,ha='center')
 for protein,y in [('TCF7L2',top-3),('GAPDH',top-9)]:
  a=axmm(x,y,width,4.6);a.imshow(crops[c,protein],aspect='auto',interpolation='none');a.axis('off');txt(x-1,y+2.3,protein,5.2,ha='right')
def wells(c,x,y,size=13,gap=1.2,rowgap=4,labeloffset=1):
 for r,(g,ps) in enumerate(colony_images[c].items()):
  yy=y-r*(size+rowgap);txt(x-labeloffset,yy+size/2,'shNC' if g=='NC' else 'shTCF7L2 '+g,5.5,ha='right')
  for j,p in enumerate(ps):
   a=axmm(x+j*(size+gap),yy,size,size);im=Image.open(p)
   if im.mode=='RGBA':
    bg=Image.new('RGBA',im.size,'white');bg.alpha_composite(im);im=bg.convert('RGB')
   a.imshow(im,interpolation='none');a.axis('off')

# S4: six cell models, original continuous WB lanes and qPCR/protein quantification.
start(180,158)
for i,c in enumerate(ORDER):
 x=[5,64,123][i%3];top=150 if i<3 else 76
 title('ABCDEF'[i],x,top,c)
 blots(c,x+12,top-12,42)
 txt(x+17,top-28,'mRNA',6.5,'bold',ha='center');txt(x+43,top-28,'Protein',6.5,'bold',ha='center')
 a=axmm(x+7,top-61,19,29);barplot(a,qpcr[c],c,'qPCR','Relative mRNA',ylim=2.05)
 a=axmm(x+33,top-61,19,29);barplot(a,wb[c],c,'WB','Relative protein',ylim=2.05)
txt(5,4,'NC = shNC; # = shTCF7L2 construct. Comparisons are against matched shNC.',5.8,color=GRAY)
save('FigS4')

# S5: control-cell growth above all available colony images and quantification.
start(180,146)
for i,c in enumerate(['LS174T','RKO','NCM460']):
 x=[15,73,131][i];title('ABC'[i],x-10,140,c)
 a=axmm(x,99,41,34);clean(a)
 for g,arr in growth[c].items():
  col=palette[g]
  a.errorbar(times,arr.mean(axis=1),yerr=arr.std(axis=1,ddof=1),color=col,marker='o',markersize=2.5,lw=1,capsize=1.6,label='shNC' if g=='NC' else 'shTCF7L2 '+g+'  ns')
 a.set_xticks(times);a.set_xlim(7,101);a.set_ylim(0,2);a.set_yticks([0,.5,1,1.5,2]);a.set_xlabel('Time (h)',fontsize=6.5)
 if i==0:a.set_ylabel('SRB absorbance (575 nm)',fontsize=6.5)
 a.legend(frameon=False,fontsize=5.5,loc='upper left',handlelength=1.3)
 title('DEF'[i],x-10,86.5,c+' colonies')
 # short row labels prevent collisions with the preceding column
 for r,(g,ps) in enumerate(colony_images[c].items()):
  yy=67-r*17;txt(x-1.2,yy+6.5,'NC' if g=='NC' else g,6,ha='right')
  for j,p in enumerate(ps):
   a=axmm(x+j*14,yy,13,13);a.imshow(Image.open(p),interpolation='none');a.axis('off')
 a=axmm(x+4,12,33,29);barplot(a,colony[c],c,'colony','Colony number',ylim=800 if c=='LS174T' else 550)
txt(5,3.5,'Growth-curve annotations reproduce the supplied thesis comparisons; colony brackets compare each KD with shNC.',5.6,color=GRAY)
save('FigS5')

# S6: all mutant colony images in the top block; all counts in the bottom block.
start(180,125)
title('A',5,119,'Colony formation')
for i,c in enumerate(ORDER[:3]):
 x=[17,75,133][i];txt(x+19.5,111,c,8,'bold',ha='center')
 for r,(g,ps) in enumerate(colony_images[c].items()):
  yy=90-r*18;txt(x-1.5,yy+7,'NC' if g=='NC' else g,6,ha='right')
  for j,p in enumerate(ps):
   a=axmm(x+j*14.5,yy,14,14);a.imshow(Image.open(p),interpolation='none');a.axis('off')
title('B',5,46,'Colony quantification')
for i,c in enumerate(ORDER[:3]):
 x=[20,78,136][i];a=axmm(x,10,35,28);barplot(a,colony[c],c,'colony','Colony number',ylim={'HT29':1250,'SW620':750,'LOVO':520}[c]);a.set_title(c,fontsize=7,pad=5)
txt(5,3,'NC = shNC; # = shTCF7L2 construct. All three available colony images per condition are shown.',5.8,color=GRAY)
save('FigS6')

with (OUT/'FigS4_S6_source_data.csv').open('w',encoding='utf-8-sig',newline='') as f:
 wr=csv.writer(f);wr.writerow(['assay','cell_line','condition','observation_index','value'])
 for assay,dd in [('WB',wb),('qPCR',qpcr),('colony',colony)]:
  for c,gg in dd.items():
   for g,a in gg.items():
    for i,v in enumerate(a):wr.writerow([assay,c,g,i+1,float(v)])
panelmap={'HT29':{'qPCR':'A','WB':'B','SRB':'C','colony':'D'},'SW620':{'qPCR':'E','WB':'F','SRB':'G','colony':'H'},'LOVO':{'qPCR':'I','WB':'J','SRB':'K','colony':'L'},'LS174T':{'qPCR':'A','WB':'B','SRB':'C','colony':'D'},'NCM460':{'qPCR':'E','WB':'F','SRB':'G','colony':'H'},'RKO':{'qPCR':'I','WB':'J','SRB':'K','colony':'L'}}
with (OUT/'Significance_source_map.csv').open('w',encoding='utf-8-sig',newline='') as f:
 wr=csv.writer(f);wr.writerow(['assay','cell_line','condition','comparison','annotation','source_figure','source_panel','scope'])
 for assay,dd in SIGNIF.items():
  for c,gg in dd.items():
   for g,star in gg.items():wr.writerow([assay,c,g,g+' vs matched shNC',star,'论文结果图/'+('突变型（改）.png' if c in ORDER[:3] else '野生型（改）.png'),panelmap[c][assay],'source group-level curve annotation; not a per-time-point or 96h test' if assay=='SRB' else 'source pairwise annotation'])
(OUT/'Fig2_S4_S6_source_map.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2),encoding='utf-8')
(OUT/'build_figures.py').write_text(Path(__file__).read_text(encoding='utf-8'),encoding='utf-8')
